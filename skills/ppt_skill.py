# -*- coding: utf-8 -*-
"""
饭加鱼：PPT 文档理解技能 v1

支持：
- 读取 .pptx 演示文稿
- 提取每页文字
- 提取表格文字
- 返回页数
- 只读，不修改原文件
"""

from pathlib import Path


def _path(path: str) -> Path:
    if not path or not str(path).strip():
        raise ValueError("PPT 路径不能为空。")

    value = str(path).strip().strip('"')
    # 修复模型偶尔传入 Windows 路径末尾的反斜杠
    value = value.rstrip("\\/")

    home = Path.home()
    special_dirs = {
        "桌面": home / "Desktop",
        "desktop": home / "Desktop",
        "下载": home / "Downloads",
        "downloads": home / "Downloads",
        "文档": home / "Documents",
        "documents": home / "Documents",
    }

    if value in special_dirs:
        return special_dirs[value].resolve()

    return Path(value).expanduser().resolve()


def read_ppt(path: str, max_chars: int = 50000) -> dict:
    target = _path(path)

    if not target.exists():
        return {"success": False, "error": f"PPT 文件不存在：{target}"}

    if not target.is_file():
        return {"success": False, "error": f"这不是文件：{target}"}

    if target.suffix.lower() != ".pptx":
        return {"success": False, "error": "当前版本支持 .pptx。"}

    try:
        from pptx import Presentation
    except ImportError:
        return {
            "success": False,
            "error": "当前 Python 环境没有安装 python-pptx，请运行：python -m pip install python-pptx",
        }

    try:
        prs = Presentation(str(target))
        slides = []
        parts = []

        for slide_index, slide in enumerate(prs.slides, start=1):
            elements = []

            for shape in slide.shapes:
                # 普通文本框/标题
                if hasattr(shape, "text") and shape.text:
                    text = shape.text.strip()
                    if text:
                        elements.append(text)

                # 表格
                if getattr(shape, "has_table", False):
                    rows = []
                    for row in shape.table.rows:
                        rows.append(" | ".join(cell.text.strip() for cell in row.cells))
                    table_text = "\n".join(rows).strip()
                    if table_text:
                        elements.append("[表格]\n" + table_text)

            slide_text = "\n".join(elements).strip()

            slides.append({
                "slide": slide_index,
                "text": slide_text,
                "char_count": len(slide_text),
            })

            if slide_text:
                parts.append(f"[第 {slide_index} 页]\n{slide_text}")

        content = "\n\n".join(parts).strip()

        return {
            "success": True,
            "path": str(target),
            "name": target.name,
            "type": "pptx",
            "slide_count": len(prs.slides),
            "slides": slides,
            "content": content[:max_chars],
            "truncated": len(content) > max_chars,
        }

    except Exception as e:
        return {
            "success": False,
            "error": f"读取 PPT 失败：{e}",
        }
