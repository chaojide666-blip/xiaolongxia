# -*- coding: utf-8 -*-
"""
饭加鱼：文档理解技能 v1

支持：
- 读取 TXT / MD / CSV / LOG / JSON 等文本文件
- 读取 DOCX 正文
- 提取 DOCX 表格内容
- 获取文档基本信息

安全：只读，不修改原文件。
"""

from pathlib import Path
import json

try:
    from docx import Document
except ImportError:
    Document = None


TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".log", ".json", ".xml",
    ".yaml", ".yml", ".ini", ".cfg", ".py", ".js", ".ts", ".html",
    ".css", ".sql"
}


def _path(path: str) -> Path:
    if not path or not str(path).strip():
        raise ValueError("文件路径不能为空。")

    value = str(path).strip()
    home = Path.home()
    special_dirs = {
        "桌面": home / "Desktop",
        "desktop": home / "Desktop",
        "下载": home / "Downloads",
        "downloads": home / "Downloads",
        "文档": home / "Documents",
        "documents": home / "Documents",
    }

    if value in special_dirs:
        return special_dirs[value].resolve()

    return Path(value).expanduser().resolve()


def _read_text_file(path: Path) -> str:
    """尽量兼容常见中文 Windows 文本编码。"""
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_docx(path: Path) -> dict:
    if Document is None:
        return {
            "success": False,
            "error": "当前 Python 环境没有安装 python-docx，请先运行：python -m pip install python-docx"
        }

    document = Document(str(path))

    paragraphs = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(text)

    tables = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({
            "table_index": table_index,
            "rows": rows,
        })

    return {
        "success": True,
        "path": str(path),
        "name": path.name,
        "type": "docx",
        "paragraphs": paragraphs,
        "paragraph_count": len(paragraphs),
        "tables": tables,
        "table_count": len(tables),
    }


def read_document(path: str, max_chars: int = 30000) -> dict:
    """读取文档正文内容，用于饭加鱼理解、总结、提取信息。"""
    target = _path(path)

    if not target.exists():
        return {"success": False, "error": f"文件不存在：{target}"}
    if not target.is_file():
        return {"success": False, "error": f"这不是文件：{target}"}

    suffix = target.suffix.lower()

    try:
        if suffix == ".docx":
            result = _read_docx(target)
            if not result.get("success"):
                return result

            parts = []
            if result["paragraphs"]:
                parts.append("\n".join(result["paragraphs"]))

            for table in result["tables"]:
                parts.append(f"\n[表格 {table['table_index']}]\n")
                for row in table["rows"]:
                    parts.append(" | ".join(row))

            content = "\n".join(parts).strip()
            result["content"] = content[:max_chars]
            result["truncated"] = len(content) > max_chars
            return result

        if suffix in TEXT_EXTENSIONS:
            content = _read_text_file(target)
            return {
                "success": True,
                "path": str(target),
                "name": target.name,
                "type": suffix.lstrip(".") or "text",
                "content": content[:max_chars],
                "truncated": len(content) > max_chars,
                "char_count": len(content),
            }

        return {
            "success": False,
            "error": (
                f"暂时不支持直接读取 {suffix or '无扩展名'} 文件。"
                "当前优先支持 DOCX 和常见文本文件。"
            ),
        }

    except Exception as e:
        return {
            "success": False,
            "error": f"读取文档失败：{e}",
        }
