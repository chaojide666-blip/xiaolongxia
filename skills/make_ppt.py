# -*- coding: utf-8 -*-
"""
饭加鱼：PPT 生成技能（Claude 式 · JSON 驱动版）

接收结构化 JSON，自动选择版式和配色，生成 PPTX。
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
import re
import time
import os
import json
from datetime import datetime


# =========================================================
# 主题配色
# =========================================================

THEMES = {
    "business": {
        "primary": RGBColor(0x1F, 0x4E, 0x79),
        "secondary": RGBColor(0x2E, 0x74, 0xB5),
        "accent": RGBColor(0xFF, 0xC0, 0x00),
        "text": RGBColor(0x33, 0x33, 0x33),
        "light": RGBColor(0xF2, 0xF2, 0xF2),
        "white": RGBColor(0xFF, 0xFF, 0xFF),
    },
    "simple": {
        "primary": RGBColor(0x40, 0x40, 0x40),
        "secondary": RGBColor(0x80, 0x80, 0x80),
        "accent": RGBColor(0xE0, 0xE0, 0xE0),
        "text": RGBColor(0x33, 0x33, 0x33),
        "light": RGBColor(0xF7, 0xF7, 0xF7),
        "white": RGBColor(0xFF, 0xFF, 0xFF),
    },
    "tech": {
        "primary": RGBColor(0x10, 0xA3, 0x7F),
        "secondary": RGBColor(0x0E, 0x8F, 0x6F),
        "accent": RGBColor(0x00, 0xD4, 0xAA),
        "text": RGBColor(0x20, 0x21, 0x23),
        "light": RGBColor(0xEC, 0xFD, 0xF5),
        "white": RGBColor(0xFF, 0xFF, 0xFF),
    },
    "warm": {
        "primary": RGBColor(0xFF, 0x98, 0x00),
        "secondary": RGBColor(0xFF, 0x70, 0x43),
        "accent": RGBColor(0xFF, 0xC1, 0x07),
        "text": RGBColor(0x33, 0x33, 0x33),
        "light": RGBColor(0xFF, 0xF8, 0xE7),
        "white": RGBColor(0xFF, 0xFF, 0xFF),
    },
}


# =========================================================
# 基础工具
# =========================================================

def _add_text(slide, left, top, width, height, text, font_size=18,
              bold=False, color=None, align=PP_ALIGN.LEFT, font_name="微软雅黑"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(text)
    p.alignment = align
    for run in p.runs:
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.name = font_name
        if color:
            run.font.color.rgb = color
    return txBox


def _add_bullet_list(slide, left, top, width, height, items, font_size=16,
                     color=None):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = "• " + str(item)
        p.space_after = Pt(8)
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.name = "微软雅黑"
            if color:
                run.font.color.rgb = color
    return txBox


def _add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def _add_table(slide, left, top, width, height, rows_data, theme):
    rows = len(rows_data)
    cols = len(rows_data[0]) if rows_data else 0
    if rows == 0 or cols == 0:
        return

    table_shape = slide.shapes.add_table(
        rows, cols, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    table = table_shape.table

    for r_idx, row in enumerate(rows_data):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = str(value)

            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = theme["primary"]
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = theme["white"]
                        run.font.size = Pt(13)
                        run.font.name = "微软雅黑"
            else:
                if r_idx % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = theme["light"]
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = theme["text"]
                        run.font.size = Pt(12)
                        run.font.name = "微软雅黑"


def _add_chart_image(slide, chart_data, chart_type, theme,
                     left=1, top=1.8, width=8, height=4.8):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from matplotlib import font_manager

        try:
            font_manager.fontManager.addfont("C:/Windows/Fonts/msyh.ttc")
            plt.rcParams["font.sans-serif"] = ["Microsoft YaHei"]
        except Exception:
            plt.rcParams["font.sans-serif"] = ["SimHei"]
        plt.rcParams["axes.unicode_minus"] = False

        labels = chart_data.get("labels", [])
        values = chart_data.get("values", [])
        title = chart_data.get("title", "")

        fig, ax = plt.subplots(figsize=(8, 4))
        if chart_type == "bar":
            ax.bar(labels, values, color="#2E74B5")
        elif chart_type == "line":
            ax.plot(labels, values, marker="o", color="#2E74B5")
        elif chart_type == "pie":
            ax.pie(values, labels=labels, autopct="%1.1f%%")

        if title:
            ax.set_title(title)
        plt.tight_layout()

        img_path = Path(__file__).parent / f"_ppt_chart_{int(time.time())}.png"
        plt.savefig(str(img_path), dpi=120)
        plt.close()

        slide.shapes.add_picture(str(img_path), Inches(left), Inches(top),
                                 width=Inches(width), height=Inches(height))

        try:
            os.remove(str(img_path))
        except Exception:
            pass

    except Exception as e:
        _add_text(slide, left, top, width, 1, f"[图表生成失败：{e}]", font_size=12)


# =========================================================
# 页面渲染
# =========================================================

def _render_cover(prs, slide_data, theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 7.5, theme["primary"])

    _add_text(slide, 0.5, 2.5, 9, 1.5, slide_data.get("title", ""),
              font_size=44, bold=True, color=theme["white"], align=PP_ALIGN.CENTER)

    if slide_data.get("subtitle"):
        _add_text(slide, 0.5, 4.2, 9, 0.8, slide_data["subtitle"],
                  font_size=18, color=theme["accent"], align=PP_ALIGN.CENTER)

    _add_text(slide, 0.5, 6.5, 9, 0.5,
              datetime.now().strftime("%Y 年 %m 月 %d 日"),
              font_size=14, color=theme["white"], align=PP_ALIGN.CENTER)


def _render_toc(prs, slide_data, theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 1.2, theme["primary"])
    _add_text(slide, 0.5, 0.3, 9, 0.8, "目录",
              font_size=32, bold=True, color=theme["white"])

    items = slide_data.get("items", [])
    for i, item in enumerate(items):
        top = 1.8 + i * 0.8
        if top > 6.8:
            break
        _add_text(slide, 1.5, top, 1, 0.6, f"{i+1:02d}",
                  font_size=24, bold=True, color=theme["accent"])
        _add_text(slide, 2.5, top + 0.1, 7, 0.6, item,
                  font_size=18, color=theme["text"])


def _render_content(prs, slide_data, theme, page_num=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 1.2, theme["primary"])
    _add_text(slide, 0.5, 0.3, 9, 0.8, slide_data.get("title", ""),
              font_size=28, bold=True, color=theme["white"])

    bullets = slide_data.get("bullets", [])
    _add_bullet_list(slide, 0.8, 1.8, 8.5, 4.5, bullets,
                     font_size=18, color=theme["text"])

    _add_text(slide, 9, 7, 0.8, 0.4, str(page_num),
              font_size=12, color=theme["secondary"], align=PP_ALIGN.RIGHT)


def _render_two_columns(prs, slide_data, theme, page_num=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 1.2, theme["primary"])
    _add_text(slide, 0.5, 0.3, 9, 0.8, slide_data.get("title", ""),
              font_size=28, bold=True, color=theme["white"])

    # 左栏
    _add_rect(slide, 0.5, 1.6, 4.3, 0.5, theme["secondary"])
    _add_text(slide, 0.6, 1.65, 4, 0.4, slide_data.get("left_title", "左栏"),
              font_size=16, bold=True, color=theme["white"])
    _add_bullet_list(slide, 0.6, 2.3, 4.2, 4.5,
                     slide_data.get("left_items", []),
                     font_size=14, color=theme["text"])

    # 右栏
    _add_rect(slide, 5.2, 1.6, 4.3, 0.5, theme["secondary"])
    _add_text(slide, 5.3, 1.65, 4, 0.4, slide_data.get("right_title", "右栏"),
              font_size=16, bold=True, color=theme["white"])
    _add_bullet_list(slide, 5.3, 2.3, 4.2, 4.5,
                     slide_data.get("right_items", []),
                     font_size=14, color=theme["text"])

    _add_text(slide, 9, 7, 0.8, 0.4, str(page_num),
              font_size=12, color=theme["secondary"], align=PP_ALIGN.RIGHT)


def _render_table(prs, slide_data, theme, page_num=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 1.2, theme["primary"])
    _add_text(slide, 0.5, 0.3, 9, 0.8, slide_data.get("title", ""),
              font_size=28, bold=True, color=theme["white"])

    rows = slide_data.get("rows", [])
    if rows:
        _add_table(slide, 0.8, 1.8, 8.4, 4.5, rows, theme)

    _add_text(slide, 9, 7, 0.8, 0.4, str(page_num),
              font_size=12, color=theme["secondary"], align=PP_ALIGN.RIGHT)


def _render_chart(prs, slide_data, theme, page_num=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 1.2, theme["primary"])
    _add_text(slide, 0.5, 0.3, 9, 0.8, slide_data.get("title", ""),
              font_size=28, bold=True, color=theme["white"])

    chart_data = slide_data.get("chart_data", {})
    chart_type = slide_data.get("chart_type", "bar")
    _add_chart_image(slide, chart_data, chart_type, theme)

    _add_text(slide, 9, 7, 0.8, 0.4, str(page_num),
              font_size=12, color=theme["secondary"], align=PP_ALIGN.RIGHT)


def _render_data_cards(prs, slide_data, theme, page_num=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 1.2, theme["primary"])
    _add_text(slide, 0.5, 0.3, 9, 0.8, slide_data.get("title", ""),
              font_size=28, bold=True, color=theme["white"])

    cards = slide_data.get("cards", [])
    n = len(cards)
    if n == 0:
        return

    card_width = 8 / n
    for i, card in enumerate(cards):
        left = 1 + i * card_width
        _add_rect(slide, left, 2.2, card_width - 0.3, 3, theme["light"])

        _add_text(slide, left, 2.6, card_width - 0.3, 1,
                  card.get("value", ""),
                  font_size=36, bold=True, color=theme["primary"],
                  align=PP_ALIGN.CENTER)

        _add_text(slide, left, 3.8, card_width - 0.3, 0.8,
                  card.get("label", ""),
                  font_size=14, color=theme["text"],
                  align=PP_ALIGN.CENTER)

    _add_text(slide, 9, 7, 0.8, 0.4, str(page_num),
              font_size=12, color=theme["secondary"], align=PP_ALIGN.RIGHT)


def _render_image(prs, slide_data, theme, page_num=1):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 1.2, theme["primary"])
    _add_text(slide, 0.5, 0.3, 9, 0.8, slide_data.get("title", ""),
              font_size=28, bold=True, color=theme["white"])

    img_path = slide_data.get("image_path", "")
    if img_path and os.path.exists(img_path):
        try:
            slide.shapes.add_picture(img_path, Inches(1.5), Inches(1.8),
                                     width=Inches(7))
        except Exception as e:
            _add_text(slide, 1, 3, 8, 1, f"[图片加载失败：{e}]", font_size=14)

    if slide_data.get("caption"):
        _add_text(slide, 0.5, 6.5, 9, 0.5, slide_data["caption"],
                  font_size=12, color=theme["secondary"], align=PP_ALIGN.CENTER)

    _add_text(slide, 9, 7, 0.8, 0.4, str(page_num),
              font_size=12, color=theme["secondary"], align=PP_ALIGN.RIGHT)


def _render_end(prs, slide_data, theme):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, 10, 7.5, theme["primary"])

    _add_text(slide, 0.5, 3, 9, 1.5,
              slide_data.get("text", "谢谢观看"),
              font_size=48, bold=True, color=theme["white"],
              align=PP_ALIGN.CENTER)

    _add_text(slide, 0.5, 4.8, 9, 0.6, "饭加鱼 · 生成",
              font_size=16, color=theme["accent"],
              align=PP_ALIGN.CENTER)


# =========================================================
# 主函数
# =========================================================

RENDERERS = {
    "cover": _render_cover,
    "toc": _render_toc,
    "content": _render_content,
    "two_columns": _render_two_columns,
    "table": _render_table,
    "chart": _render_chart,
    "data_cards": _render_data_cards,
    "image": _render_image,
    "end": _render_end,
}


def make_ppt(topic, pages=6, structure_json=None):
    """
    生成 PPT。
    topic: 主题（用于文件名）
    pages: 页数（仅在 structure_json 为空时用）
    structure_json: 结构化 JSON（字符串或字典）
    """
    project_dir = Path(__file__).parent
    output_dir = project_dir / "生成的PPT"
    output_dir.mkdir(exist_ok=True)

    safe_topic = re.sub(r'[\\/:*?"<>|]', "", topic)
    file_path = output_dir / f"{safe_topic}_{int(time.time())}.pptx"

    # 解析 JSON
    structure = None
    if structure_json:
        try:
            if isinstance(structure_json, str):
                structure = json.loads(structure_json)
            else:
                structure = structure_json
        except Exception as e:
            print(f"JSON 解析失败：{e}")

    # 如果没有 JSON，用兜底模板
    if not structure:
        structure = _fallback_structure(topic, pages)

    theme_name = structure.get("theme", "business")
    theme = THEMES.get(theme_name, THEMES["business"])
    slides = structure.get("slides", [])

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    page_num = 1
    for slide_data in slides:
        slide_type = slide_data.get("type", "content")
        renderer = RENDERERS.get(slide_type)

        if not renderer:
            continue

        try:
            if slide_type in ("cover", "toc", "end"):
                renderer(prs, slide_data, theme)
            else:
                renderer(prs, slide_data, theme, page_num)
                page_num += 1
        except Exception as e:
            print(f"渲染 {slide_type} 页失败：{e}")
            continue

    prs.save(str(file_path))
    print("PPT生成成功:", file_path)
    return str(file_path)


def _fallback_structure(topic, pages):
    """没有 JSON 时的兜底结构。"""
    content_pages = max(1, pages - 3)
    slides = [
        {"type": "cover", "title": topic, "subtitle": "饭加鱼 出品"},
        {"type": "toc", "items": [f"第 {i+1} 部分" for i in range(content_pages)]},
    ]
    for i in range(content_pages):
        slides.append({
            "type": "content",
            "title": f"第 {i+1} 部分",
            "bullets": ["要点一", "要点二", "要点三"],
        })
    slides.append({"type": "end"})
    return {"topic": topic, "theme": "business", "slides": slides}
